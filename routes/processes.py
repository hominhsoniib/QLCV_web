from fastapi import APIRouter, Depends, Request, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from database.connection import get_db
from services.auth_service import AuthService
from services.ai_service import AIService
from models.models import Process, ProcessStep, Department, Employee
import json
import io
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

router = APIRouter()
templates = Jinja2Templates(directory="templates")

# Helper to style Word table cells
def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

@router.get("/processes", response_class=HTMLResponse)
def get_processes_page(request: Request, db: Session = Depends(get_db)):
    """Renders the processes management page. Requires login."""
    session = AuthService.get_session(request)
    if not session:
        return RedirectResponse(url="/login")
    
    departments = db.query(Department).all()
    processes = db.query(Process).all()
    return templates.TemplateResponse(request, "processes.html", {
        "user": session,
        "departments": departments,
        "processes": processes
    })

@router.get("/api/processes")
def api_list_processes(request: Request, db: Session = Depends(get_db)):
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
    
    processes = db.query(Process).all()
    data = []
    for p in processes:
        data.append({
            "id": p.id,
            "dept_code": p.dept_code,
            "code": p.code,
            "name": p.name,
            "type": p.type,
            "status": p.status,
            "version": p.version
        })
    return {"success": True, "data": data}

@router.get("/api/processes/{id}")
def api_get_process_detail(id: int, request: Request, db: Session = Depends(get_db)):
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
    
    process = db.query(Process).filter(Process.id == id).first()
    if not process:
        return {"success": False, "message": "Không tìm thấy quy trình!"}
        
    steps = db.query(ProcessStep).filter(ProcessStep.process_id == id).order_by(ProcessStep.step_no).all()
    
    steps_data = []
    for s in steps:
        steps_data.append({
            "step_no": s.step_no,
            "name": s.name,
            "owner": s.owner,
            "input": s.input,
            "output": s.output,
            "sla": s.sla
        })
        
    # Safely parse JSON strings
    try:
        checklist = json.loads(process.checklist)
    except:
        checklist = []
        
    try:
        kpis = json.loads(process.kpi_suggestions)
    except:
        kpis = []
        
    return {
        "success": True,
        "data": {
            "id": process.id,
            "dept_code": process.dept_code,
            "code": process.code,
            "name": process.name,
            "type": process.type,
            "objective": process.objective,
            "scope": process.scope,
            "flowchart_mermaid": process.flowchart_mermaid,
            "checklist": checklist,
            "kpi_suggestions": kpis,
            "status": process.status,
            "version": process.version,
            "steps": steps_data
        }
    }

@router.post("/api/processes/generate")
def api_generate_process_draft(request: Request, payload: dict, db: Session = Depends(get_db)):
    """Triggers AI generation for the process structure."""
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
        
    prompt = payload.get("prompt", "").strip()
    if not prompt:
        return {"success": False, "message": "Nội dung yêu cầu không được trống!"}
        
    try:
        draft = AIService.generate_process(db, prompt)
        return {"success": True, "data": draft}
    except Exception as e:
        return {"success": False, "message": f"AI Generation Error: {str(e)}"}

@router.post("/api/processes")
def api_create_process(request: Request, payload: dict, db: Session = Depends(get_db)):
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
    if session["role"] not in ["ADMIN", "CEO", "MANAGER"]:
        return {"success": False, "message": "Chỉ có Admin/CEO/Manager mới có quyền tạo quy trình!"}
        
    try:
        # Check if code already exists
        existing = db.query(Process).filter(Process.code == payload.get("code")).first()
        if existing:
            return {"success": False, "message": f"Mã quy trình {payload.get('code')} đã tồn tại!"}
            
        process = Process(
            dept_code=payload.get("dept_code"),
            code=payload.get("code"),
            name=payload.get("name"),
            type=payload.get("type", "Quy trình ISO"),
            objective=payload.get("objective", ""),
            scope=payload.get("scope", ""),
            flowchart_mermaid=payload.get("flowchart_mermaid", ""),
            checklist=json.dumps(payload.get("checklist", [])),
            kpi_suggestions=json.dumps(payload.get("kpi_suggestions", [])),
            status=payload.get("status", "draft"),
            version=1
        )
        db.add(process)
        db.commit()
        db.refresh(process)
        
        # Save steps
        steps = payload.get("steps", [])
        for s in steps:
            step = ProcessStep(
                process_id=process.id,
                step_no=s.get("step_no"),
                name=s.get("name"),
                owner=s.get("owner"),
                input=s.get("input"),
                output=s.get("output"),
                sla=s.get("sla")
            )
            db.add(step)
            
        db.commit()
        return {"success": True, "message": "✅ Lưu quy trình mới thành công!", "id": process.id}
    except Exception as e:
        db.rollback()
        return {"success": False, "message": f"Lỗi tạo quy trình: {str(e)}"}

@router.put("/api/processes/{id}")
def api_update_process(id: int, request: Request, payload: dict, db: Session = Depends(get_db)):
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
    if session["role"] not in ["ADMIN", "CEO", "MANAGER"]:
        return {"success": False, "message": "Chỉ có Admin/CEO/Manager mới có quyền cập nhật quy trình!"}
        
    try:
        process = db.query(Process).filter(Process.id == id).first()
        if not process:
            return {"success": False, "message": "Không tìm thấy quy trình!"}
            
        process.dept_code = payload.get("dept_code")
        process.code = payload.get("code")
        process.name = payload.get("name")
        process.type = payload.get("type")
        process.objective = payload.get("objective")
        process.scope = payload.get("scope")
        process.flowchart_mermaid = payload.get("flowchart_mermaid")
        process.checklist = json.dumps(payload.get("checklist", []))
        process.kpi_suggestions = json.dumps(payload.get("kpi_suggestions", []))
        process.status = payload.get("status")
        process.version = process.version + 1
        
        # Delete old steps
        db.query(ProcessStep).filter(ProcessStep.process_id == id).delete()
        
        # Add new steps
        steps = payload.get("steps", [])
        for s in steps:
            step = ProcessStep(
                process_id=id,
                step_no=s.get("step_no"),
                name=s.get("name"),
                owner=s.get("owner"),
                input=s.get("input"),
                output=s.get("output"),
                sla=s.get("sla")
            )
            db.add(step)
            
        db.commit()
        return {"success": True, "message": "✅ Cập nhật quy trình thành công!"}
    except Exception as e:
        db.rollback()
        return {"success": False, "message": f"Lỗi cập nhật quy trình: {str(e)}"}

@router.delete("/api/processes/{id}")
def api_delete_process(id: int, request: Request, db: Session = Depends(get_db)):
    session = AuthService.get_session(request)
    if not session:
        return {"success": False, "message": "Chưa đăng nhập!"}
    if session["role"] not in ["ADMIN", "CEO", "MANAGER"]:
        return {"success": False, "message": "Chỉ có Admin/CEO/Manager mới có quyền xóa quy trình!"}
        
    try:
        db.query(ProcessStep).filter(ProcessStep.process_id == id).delete()
        db.query(Process).filter(Process.id == id).delete()
        db.commit()
        return {"success": True, "message": "✅ Xóa quy trình thành công!"}
    except Exception as e:
        db.rollback()
        return {"success": False, "message": f"Lỗi xóa quy trình: {str(e)}"}

@router.get("/api/processes/{id}/export-word")
def api_export_process_to_word(id: int, request: Request, db: Session = Depends(get_db)):
    """Exports process and steps into a styled, professional Word docx file."""
    session = AuthService.get_session(request)
    if not session:
        return RedirectResponse(url="/login")
        
    process = db.query(Process).filter(Process.id == id).first()
    if not process:
        raise HTTPException(status_code=404, detail="Quy trình không tồn tại")
        
    steps = db.query(ProcessStep).filter(ProcessStep.process_id == id).order_by(ProcessStep.step_no).all()
    
    # Create docx Document
    doc = Document()
    
    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(process.name.upper())
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(11, 37, 69) # Dark Navy
    
    # Spacing
    doc.add_paragraph()
    
    # Section: Meta Info (General Details Table)
    table_meta = doc.add_table(rows=4, cols=2)
    table_meta.style = 'Light Shading Accent 1'
    
    meta_headers = ["MÃ QUY TRÌNH", "LOẠI QUY TRÌNH / PHIÊN BẢN", "MỤC TIÊU", "PHẠM VI ÁP DỤNG"]
    meta_values = [
        process.code,
        f"{process.type or 'ISO'} / Phiên bản {process.version}",
        process.objective or "N/A",
        process.scope or "N/A"
    ]
    
    for i in range(4):
        row = table_meta.rows[i]
        
        cell_lbl = row.cells[0]
        cell_lbl.text = meta_headers[i]
        set_cell_background(cell_lbl, "EBF4F6")
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        cell_lbl.paragraphs[0].runs[0].font.size = Pt(10)
        
        cell_val = row.cells[1]
        cell_val.text = meta_values[i]
        cell_val.paragraphs[0].runs[0].font.size = Pt(10)
        
    doc.add_paragraph()
    
    # Section Header: Flowchart
    if process.flowchart_mermaid:
        p_flow_h = doc.add_paragraph()
        run_flow_h = p_flow_h.add_run("SƠ ĐỒ QUY TRÌNH (FLOWCHART)")
        run_flow_h.font.bold = True
        run_flow_h.font.size = Pt(14)
        run_flow_h.font.color.rgb = RGBColor(11, 37, 69)
        
        try:
            import base64
            import urllib.request
            import io
            
            mermaid_bytes = process.flowchart_mermaid.strip().encode('utf-8')
            encoded = base64.urlsafe_b64encode(mermaid_bytes).decode('utf-8')
            url = f"https://mermaid.ink/img/{encoded}"
            
            req = urllib.request.Request(
                url,
                headers={'User-Agent': 'Mozilla/5.0'}
            )
            with urllib.request.urlopen(req, timeout=12) as response:
                img_data = response.read()
                img_stream = io.BytesIO(img_data)
                
                doc.add_picture(img_stream, width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        except Exception as img_err:
            print(f"[EXPORT WORD] Failed to render flowchart image: {img_err}")
            # Fallback: write the mermaid text code inside a styled paragraph
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.5)
            run_code = p_code.add_run(f"Mã sơ đồ Mermaid (Chưa tải được hình ảnh do không có mạng):\n{process.flowchart_mermaid}")
            run_code.font.name = 'Courier New'
            run_code.font.size = Pt(9.5)
            
        doc.add_paragraph()
        
    # Section Header: Detailed Steps
    p_steps_h = doc.add_paragraph()
    run_steps_h = p_steps_h.add_run("CHI TIẾT CÁC BƯỚC THỰC HIỆN")
    run_steps_h.font.bold = True
    run_steps_h.font.size = Pt(14)
    run_steps_h.font.color.rgb = RGBColor(11, 37, 69)
    
    # Steps Table
    table_steps = doc.add_table(rows=1, cols=6)
    table_steps.style = 'Table Grid'
    
    headers = ["STT", "TÊN BƯỚC CÔNG VIỆC", "BỘ PHẬN CHỦ TRÌ", "ĐẦU VÀO", "ĐẦU RA", "SLA"]
    hdr_cells = table_steps.rows[0].cells
    for i in range(6):
        hdr_cells[i].text = headers[i]
        set_cell_background(hdr_cells[i], "0B2545") # Dark Navy
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
        
    for s in steps:
        row_cells = table_steps.add_row().cells
        row_cells[0].text = str(s.step_no)
        row_cells[1].text = s.name or ""
        row_cells[2].text = s.owner or ""
        row_cells[3].text = s.input or ""
        row_cells[4].text = s.output or ""
        row_cells[5].text = s.sla or ""
        
        # Style row text size
        for cell in row_cells:
            if cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].font.size = Pt(9.5)
                
    doc.add_paragraph()
    
    # Section Header: Checklist
    try:
        checklist_items = json.loads(process.checklist)
    except:
        checklist_items = []
        
    if checklist_items:
        p_check_h = doc.add_paragraph()
        run_check_h = p_check_h.add_run("DANH MỤC KIỂM TRA (CHECKLIST)")
        run_check_h.font.bold = True
        run_check_h.font.size = Pt(14)
        run_check_h.font.color.rgb = RGBColor(11, 37, 69)
        
        for item in checklist_items:
            p_item = doc.add_paragraph(style='List Bullet')
            run_item = p_item.add_run(item)
            run_item.font.size = Pt(10)
            
        doc.add_paragraph()
        
    # Section Header: KPIs
    try:
        kpi_items = json.loads(process.kpi_suggestions)
    except:
        kpi_items = []
        
    if kpi_items:
        p_kpi_h = doc.add_paragraph()
        run_kpi_h = p_kpi_h.add_run("CHỈ SỐ ĐO LƯỜNG & ĐÁNH GIÁ (KPI)")
        run_kpi_h.font.bold = True
        run_kpi_h.font.size = Pt(14)
        run_kpi_h.font.color.rgb = RGBColor(11, 37, 69)
        
        for item in kpi_items:
            p_item = doc.add_paragraph(style='List Bullet')
            run_item = p_item.add_run(item)
            run_item.font.size = Pt(10)
            
    # Save document to memory stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    encoded_filename = quote(f"{process.name}.docx")
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=utf-8''{encoded_filename}"}
    )
